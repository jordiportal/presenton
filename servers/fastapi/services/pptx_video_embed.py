from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Iterable
from urllib.parse import unquote, urlparse
from uuid import UUID

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from utils.asset_directory_utils import resolve_app_path_to_filesystem

LOGGER = logging.getLogger(__name__)

CANVAS_WIDTH = 1280
CANVAS_HEIGHT = 720
EMBED_EXTENSIONS = {".mp4", ".m4v", ".mov"}
MIME_BY_EXTENSION = {
    ".mp4": "video/mp4",
    ".m4v": "video/mp4",
    ".mov": "video/quicktime",
}
SKIP_PROVIDERS = {"youtube", "vimeo"}


@dataclass(frozen=True)
class VideoPlacement:
    slide_index: int
    x: float
    y: float
    width: float
    height: float
    path: str
    poster_path: str | None
    mime_type: str


def collect_video_elements(value: Any) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []

    def visit(node: Any) -> None:
        if isinstance(node, list):
            for item in node:
                visit(item)
            return
        if not isinstance(node, dict):
            return
        if node.get("type") == "video":
            found.append(node)
            return
        children = node.get("children")
        if isinstance(children, list):
            visit(children)
        child = node.get("child")
        if child is not None:
            visit(child)
        elements = node.get("elements")
        if isinstance(elements, list):
            visit(elements)
        components = node.get("components")
        if isinstance(components, list):
            visit(components)

    visit(value)
    return found


def canvas_box_to_emu(
    x: float,
    y: float,
    width: float,
    height: float,
    slide_width: int,
    slide_height: int,
) -> tuple[int, int, int, int]:
    return (
        int(round(x / CANVAS_WIDTH * slide_width)),
        int(round(y / CANVAS_HEIGHT * slide_height)),
        max(1, int(round(width / CANVAS_WIDTH * slide_width))),
        max(1, int(round(height / CANVAS_HEIGHT * slide_height))),
    )


def boxes_overlap_for_replace(
    picture: tuple[int, int, int, int],
    video: tuple[int, int, int, int],
    slide_width: int,
    slide_height: int,
) -> bool:
    picture_area = max(0, picture[2]) * max(0, picture[3])
    slide_area = max(1, slide_width * slide_height)
    if picture_area <= 0 or picture_area > 0.65 * slide_area:
        return False
    inter = _intersection_area(picture, video)
    if inter <= 0:
        return False
    video_area = max(1, video[2] * video[3])
    return (inter / picture_area) >= 0.55 or (inter / video_area) >= 0.45


def resolve_local_media_path(src: str | None) -> str | None:
    if not src or not isinstance(src, str):
        return None
    candidate = src.strip()
    if not candidate or candidate.startswith(("data:", "blob:", "javascript:")):
        return None
    parsed = urlparse(candidate)
    if parsed.scheme in {"http", "https"}:
        candidate = unquote(parsed.path or "")
    return resolve_app_path_to_filesystem(candidate)


def placement_from_element(slide_index: int, element: dict[str, Any]) -> VideoPlacement | None:
    provider = str(element.get("provider") or "file").strip().lower()
    if provider in SKIP_PROVIDERS:
        return None
    src = element.get("src")
    if not isinstance(src, str) or not src.strip():
        return None
    path = resolve_local_media_path(src)
    if not path:
        return None
    extension = os.path.splitext(path)[1].lower()
    if extension not in EMBED_EXTENSIONS:
        return None
    position = element.get("position") if isinstance(element.get("position"), dict) else {}
    size = element.get("size") if isinstance(element.get("size"), dict) else {}
    try:
        x = float(position.get("x"))
        y = float(position.get("y"))
        width = float(size.get("width"))
        height = float(size.get("height"))
    except (TypeError, ValueError):
        return None
    if width <= 0 or height <= 0:
        return None
    poster = element.get("poster")
    poster_path = (
        resolve_local_media_path(poster) if isinstance(poster, str) else None
    )
    return VideoPlacement(
        slide_index=slide_index,
        x=x,
        y=y,
        width=width,
        height=height,
        path=path,
        poster_path=poster_path,
        mime_type=MIME_BY_EXTENSION[extension],
    )


def collect_placements(slides: Iterable[Any]) -> list[VideoPlacement]:
    placements: list[VideoPlacement] = []
    for slide in slides:
        for blob in _slide_layout_blobs(slide):
            for element in collect_video_elements(blob):
                placement = placement_from_element(slide.index, element)
                if placement:
                    placements.append(placement)
    return placements


def embed_videos_in_pptx(pptx_path: str, placements: list[VideoPlacement]) -> int:
    if not placements or not os.path.isfile(pptx_path):
        return 0
    presentation = Presentation(pptx_path)
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    by_index: dict[int, list[VideoPlacement]] = {}
    for placement in placements:
        by_index.setdefault(placement.slide_index, []).append(placement)

    embedded = 0
    for slide_index, slide in enumerate(presentation.slides):
        for placement in by_index.get(slide_index, []):
            if _embed_on_slide(
                slide,
                placement,
                slide_width,
                slide_height,
            ):
                embedded += 1
    if embedded:
        presentation.save(pptx_path)
    return embedded


async def embed_presentation_videos(
    pptx_path: str,
    presentation_id: UUID,
) -> int:
    from sqlmodel import select

    from models.sql.slide import SlideModel
    from services.database import async_session_maker

    async with async_session_maker() as session:
        slides = list(
            await session.scalars(
                select(SlideModel)
                .where(SlideModel.presentation == presentation_id)
                .order_by(SlideModel.index)
            )
        )
    placements = collect_placements(slides)
    if not placements:
        return 0
    return embed_videos_in_pptx(pptx_path, placements)


def _slide_layout_blobs(slide: Any) -> list[dict[str, Any]]:
    blobs: list[dict[str, Any]] = []
    if isinstance(slide.ui, dict):
        blobs.append(slide.ui)
    content = slide.content
    if isinstance(content, dict) and (
        isinstance(content.get("elements"), list)
        or isinstance(content.get("components"), list)
    ):
        blobs.append(content)
    return blobs


def _intersection_area(
    first: tuple[int, int, int, int],
    second: tuple[int, int, int, int],
) -> int:
    ax, ay, aw, ah = first
    bx, by, bw, bh = second
    left = max(ax, bx)
    top = max(ay, by)
    right = min(ax + aw, bx + bw)
    bottom = min(ay + ah, by + bh)
    if right <= left or bottom <= top:
        return 0
    return (right - left) * (bottom - top)


def _embed_on_slide(
    slide: Any,
    placement: VideoPlacement,
    slide_width: int,
    slide_height: int,
) -> bool:
    box = canvas_box_to_emu(
        placement.x,
        placement.y,
        placement.width,
        placement.height,
        slide_width,
        slide_height,
    )
    placeholders = _overlapping_pictures(slide, box, slide_width, slide_height)
    poster = _poster_stream(placeholders, placement)
    try:
        slide.shapes.add_movie(
            placement.path,
            box[0],
            box[1],
            box[2],
            box[3],
            poster_frame_image=poster,
            mime_type=placement.mime_type,
        )
    except Exception:
        LOGGER.exception(
            "Failed to embed video %s on slide %s",
            placement.path,
            placement.slide_index,
        )
        return False
    for shape in placeholders:
        element = shape._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)
    return True


def _overlapping_pictures(
    slide: Any,
    box: tuple[int, int, int, int],
    slide_width: int,
    slide_height: int,
) -> list[Any]:
    matches: list[Any] = []
    for shape in slide.shapes:
        if getattr(shape, "shape_type", None) != MSO_SHAPE_TYPE.PICTURE:
            continue
        picture_box = (
            int(shape.left),
            int(shape.top),
            int(shape.width),
            int(shape.height),
        )
        if boxes_overlap_for_replace(picture_box, box, slide_width, slide_height):
            matches.append(shape)
    return matches


def _poster_stream(
    placeholders: list[Any],
    placement: VideoPlacement,
) -> BytesIO | str | None:
    if placement.poster_path and os.path.isfile(placement.poster_path):
        return placement.poster_path
    for shape in placeholders:
        try:
            return BytesIO(shape.image.blob)
        except Exception:
            continue
    return None
