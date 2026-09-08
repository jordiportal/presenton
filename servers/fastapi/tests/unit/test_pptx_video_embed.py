from types import SimpleNamespace

from pptx import Presentation
from pptx.util import Inches
from PIL import Image

from services.pptx_video_embed import (
    VideoPlacement,
    boxes_overlap_for_replace,
    canvas_box_to_emu,
    collect_video_elements,
    embed_videos_in_pptx,
    placement_from_element,
)


def test_collect_video_elements_walks_nested_layout():
    ui = {
        "components": [
            {
                "elements": [
                    {
                        "type": "group",
                        "children": [
                            {
                                "type": "video",
                                "src": "/app_data/videos/clip.mp4",
                                "provider": "file",
                            }
                        ],
                    }
                ]
            }
        ]
    }
    found = collect_video_elements(ui)
    assert len(found) == 1
    assert found[0]["src"].endswith("clip.mp4")


def test_skips_youtube_and_keeps_mp4():
    youtube = placement_from_element(
        0,
        {
            "type": "video",
            "provider": "youtube",
            "src": "https://www.youtube.com/embed/dQw4w9wgGcQ",
            "position": {"x": 10, "y": 10},
            "size": {"width": 100, "height": 100},
        },
    )
    assert youtube is None


def test_canvas_box_to_emu_scales_16_9():
    slide_w = int(Inches(13.333333))
    slide_h = int(Inches(7.5))
    left, top, width, height = canvas_box_to_emu(0, 0, 1280, 720, slide_w, slide_h)
    assert left == 0
    assert top == 0
    assert width == slide_w
    assert height == slide_h


def test_overlap_ignores_full_slide_picture():
    slide = (0, 0, 1280, 720)
    video = (280, 158, 720, 405)
    assert boxes_overlap_for_replace(slide, video, 1280, 720) is False
    poster = (280, 158, 720, 405)
    assert boxes_overlap_for_replace(poster, video, 1280, 720) is True


def test_embed_videos_in_pptx_adds_movie_part(tmp_path, monkeypatch):
    pptx_path = tmp_path / "deck.pptx"
    video_path = tmp_path / "clip.mp4"
    video_path.write_bytes(b"\x00\x00\x00\x18ftypisom" + b"\x00" * 64)
    poster = tmp_path / "poster.png"
    Image.new("RGB", (8, 8), color=(10, 20, 30)).save(poster, format="PNG")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(pptx_path)

    added = []

    def fake_add_movie(self, path, left, top, width, height, poster_frame_image=None, mime_type="video/mp4"):
        added.append(
            SimpleNamespace(
                path=path,
                left=left,
                top=top,
                width=width,
                height=height,
                mime_type=mime_type,
            )
        )
        return SimpleNamespace()

    monkeypatch.setattr(
        "pptx.slide.SlideShapes.add_movie",
        fake_add_movie,
        raising=False,
    )
    monkeypatch.setattr(
        "pptx.shapes.shapetree.SlideShapes.add_movie",
        fake_add_movie,
        raising=False,
    )

    count = embed_videos_in_pptx(
        str(pptx_path),
        [
            VideoPlacement(
                slide_index=0,
                x=280,
                y=158,
                width=720,
                height=405,
                path=str(video_path),
                poster_path=str(poster),
                mime_type="video/mp4",
            )
        ],
    )
    assert count == 1
    assert added
    assert added[0].path == str(video_path)
    assert added[0].mime_type == "video/mp4"
    assert added[0].left > 0
    assert added[0].width > 0
